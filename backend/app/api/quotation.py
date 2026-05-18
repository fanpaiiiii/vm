"""报价单Excel导出API - 支持logo图片嵌入"""
import io
import base64
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.drawing.image import Image as XlImage
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/api/quotation", tags=["报价单"])


class ProductItem(BaseModel):
    size: str = ""
    qty: int = 1
    unitPrice: float = 0
    amount: float = 0


class ProductGroup(BaseModel):
    productNo: str = ""
    spec: str = ""
    items: list[ProductItem] = []
    imageUrl: Optional[str] = None  # base64 data URL


class ShippingInfo(BaseModel):
    method: str = "Shipping cost"
    cost: float = 0


class CompanyInfo(BaseModel):
    name: str = ""
    address: str = ""
    tel: str = ""
    logoUrl: Optional[str] = None  # base64 data URL


class QuotationRequest(BaseModel):
    company: CompanyInfo = CompanyInfo()
    buyerName: str = ""
    attn: str = ""
    quotationNo: str = ""
    dates: str = ""
    validDates: str = ""
    currency: str = "USD"
    tradeTerms: str = "DDP"
    groups: list[ProductGroup] = []
    shipping: ShippingInfo = ShippingInfo()


def _data_url_to_image(data_url: str) -> Optional[XlImage]:
    """Convert base64 data URL to openpyxl Image, return None if invalid."""
    if not data_url or not data_url.startswith("data:image/"):
        return None
    try:
        # Extract base64 part
        header, b64data = data_url.split(",", 1)
        img_bytes = base64.b64decode(b64data)
        img_stream = io.BytesIO(img_bytes)
        img = XlImage(img_stream)
        return img
    except Exception:
        return None


@router.post("/export/excel")
async def export_quotation_excel(req: QuotationRequest):
    """导出报价单为Excel，支持logo和产品图片嵌入"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"

    # Column widths: A=5, B=14, C=16, D=14, E=8, F=12, G=14
    col_widths = [5, 14, 16, 14, 8, 12, 14]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Styles
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    bold_font = Font(bold=True)
    title_font = Font(bold=True, size=16)
    company_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center')

    row = 1

    # Row 1: Company Name (B1:G1)
    ws.merge_cells('B1:G1')
    cell = ws.cell(row=1, column=2, value=req.company.name or "Company Name")
    cell.font = company_font
    cell.alignment = Alignment(horizontal='left', vertical='center')

    # Logo at A1:B5 if provided
    logo_img = _data_url_to_image(req.company.logoUrl) if req.company.logoUrl else None
    if logo_img:
        logo_img.width = 120
        logo_img.height = 60
        ws.add_image(logo_img, 'A1')

    # Row 2: Add: address
    ws.merge_cells('B2:G2')
    ws.cell(row=2, column=2, value=f"Add: {req.company.address or '-'}").font = bold_font

    # Row 3: TEL:
    ws.merge_cells('B3:G3')
    ws.cell(row=3, column=2, value=f"TEL: {req.company.tel or '-'}").font = Font(bold=True, size=10)

    # Row 5: Quotation title (B5:G5)
    ws.merge_cells('B5:G5')
    cell = ws.cell(row=5, column=2, value="Quotation")
    cell.font = title_font
    cell.alignment = center_align
    # Bottom border for title row
    for c in range(2, 8):
        ws.cell(row=5, column=c).border = Border(bottom=Side(style='medium'))

    # Row 6: To: buyer | Quotation No.
    ws.merge_cells('B6:D6')
    ws.cell(row=6, column=2, value=f"To: {req.buyerName or '________________'}").font = bold_font
    ws.merge_cells('E6:G6')
    ws.cell(row=6, column=5, value=f"Quotation No.: {req.quotationNo or '-'}").font = bold_font
    ws.cell(row=6, column=5).alignment = Alignment(horizontal='right')

    # Row 7: Attn: contact | Dates
    ws.merge_cells('B7:D7')
    ws.cell(row=7, column=2, value=f"Attn: {req.attn or '________________'}").font = bold_font
    ws.merge_cells('E7:G7')
    ws.cell(row=7, column=5, value=f"Dates: {req.dates or '-'}").font = bold_font
    ws.cell(row=7, column=5).alignment = Alignment(horizontal='right')

    # Row 8: (blank) | Valid dates
    ws.merge_cells('E8:G8')
    ws.cell(row=8, column=5, value=f"Valid dates: {req.validDates or '-'}").font = bold_font
    ws.cell(row=8, column=5).alignment = Alignment(horizontal='right')

    # Row 9: Table headers
    headers = ['No.', 'product No.', 'product picture', 'Size', 'Qty', 'unit price', 'total amount']
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=9, column=ci, value=h)
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Row 10: Trade terms
    row = 10
    ws.merge_cells(f'A{row}:G{row}')
    cell = ws.cell(row=row, column=1, value=f"{req.tradeTerms} PRICE")
    cell.font = Font(bold=True, size=12)
    cell.alignment = center_align
    cell.border = thin_border
    for c in range(2, 8):
        ws.cell(row=row, column=c).border = thin_border
    row += 1

    # Product groups
    for gi, group in enumerate(req.groups):
        start_row = row
        for ii, item in enumerate(group.items):
            ws.cell(row=row, column=4, value=item.size or '-').border = thin_border
            ws.cell(row=row, column=4).alignment = left_align
            ws.cell(row=row, column=5, value=item.qty).border = thin_border
            ws.cell(row=row, column=5).alignment = center_align
            ws.cell(row=row, column=6, value=item.unitPrice).border = thin_border
            ws.cell(row=row, column=6).alignment = right_align
            ws.cell(row=row, column=6).number_format = '#,##0.00'
            ws.cell(row=row, column=7, value=item.amount).border = thin_border
            ws.cell(row=row, column=7).alignment = right_align
            ws.cell(row=row, column=7).number_format = '#,##0.00'

            # Add borders to merged cells
            for c in [1, 2, 3]:
                ws.cell(row=row, column=c).border = thin_border

            row += 1

        end_row = row - 1
        if end_row > start_row:
            # Merge No. column
            ws.merge_cells(f'A{start_row}:A{end_row}')
            # Merge Product No. column
            ws.merge_cells(f'B{start_row}:B{end_row}')
            # Merge Picture column
            ws.merge_cells(f'C{start_row}:C{end_row}')

        # No.
        ws.cell(row=start_row, column=1, value=gi + 1).alignment = center_align
        ws.cell(row=start_row, column=1).border = thin_border
        # Product No.
        ws.cell(row=start_row, column=2, value=group.productNo or '-').alignment = center_align
        ws.cell(row=start_row, column=2).border = thin_border
        # Picture cell
        pic_cell = ws.cell(row=start_row, column=3)
        pic_cell.alignment = center_align
        pic_cell.border = thin_border

        # Embed product image if provided
        prod_img = _data_url_to_image(group.imageUrl) if group.imageUrl else None
        if prod_img:
            prod_img.width = 100
            prod_img.height = 80
            cell_ref = f'C{start_row}'
            ws.add_image(prod_img, cell_ref)
        else:
            pic_cell.value = "[Picture]"

    # Shipping row
    ws.cell(row=row, column=1).border = thin_border
    ws.merge_cells(f'B{row}:F{row}')
    ws.cell(row=row, column=2, value=req.shipping.method or 'Shipping cost').font = Font(bold=True, size=11)
    ws.cell(row=row, column=2).alignment = center_align
    ws.cell(row=row, column=2).border = thin_border
    for c in range(3, 7):
        ws.cell(row=row, column=c).border = thin_border
    ws.cell(row=row, column=7, value=req.shipping.cost).border = thin_border
    ws.cell(row=row, column=7).alignment = right_align
    ws.cell(row=row, column=7).number_format = '#,##0.00'
    ws.cell(row=row, column=7).font = bold_font
    row += 1

    # Total row
    ws.merge_cells(f'A{row}:F{row}')
    cell = ws.cell(row=row, column=1, value=f"TOTAL AMOUNT({req.tradeTerms}):")
    cell.font = Font(bold=True, size=13)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = Border(top=Side(style='double'), bottom=Side(style='double'))
    for c in range(2, 7):
        ws.cell(row=row, column=c).border = Border(top=Side(style='double'), bottom=Side(style='double'))

    # Calculate grand total
    grand_total = sum(
        item.amount for g in req.groups for item in g.items
    ) + req.shipping.cost

    currency_symbols = {'USD': '$', 'EUR': '€', 'GBP': '£', 'CNY': '¥'}
    sym = currency_symbols.get(req.currency, '$')

    total_cell = ws.cell(row=row, column=7, value=grand_total)
    total_cell.font = Font(bold=True, size=13)
    total_cell.alignment = right_align
    total_cell.number_format = '#,##0.00'
    total_cell.border = Border(top=Side(style='double'), bottom=Side(style='double'))

    # Set row heights
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[5].height = 30
    for r in range(9, row + 1):
        ws.row_dimensions[r].height = 22

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Quotation_{req.quotationNo or 'draft'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
